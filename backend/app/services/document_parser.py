import io
import re
import html
import base64
import zipfile
import logging
import urllib.parse
from typing import Dict, Any, List, Optional

import httpx

from app.core.net import (
    ContentTooLargeError,
    UnsafeURLError,
    fetch_public_bytes,
)

logger = logging.getLogger("document_parser")

MAX_DOCUMENT_CHARS = 50000  # Safe boundary for document context injection

#: Hard ceiling on a single remote download (25MB).
MAX_REMOTE_FETCH_BYTES = 25 * 1024 * 1024

#: Ceiling on a single decoded attachment (20MB).
MAX_ATTACHMENT_BYTES = 20 * 1024 * 1024

_FETCH_HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
        "(KHTML, like Gecko) Chrome/122.0.0.0 Safari/537.36"
    ),
    "Accept": "*/*",
}

URL_PATTERN = re.compile(r"https?://[^\s<>\"'()]+", re.IGNORECASE)
_CHARSET_PATTERN = re.compile(r"charset\s*=\s*\"?([\w\-]+)", re.IGNORECASE)


def extract_urls(text: str) -> List[str]:
    """Extract all HTTP/HTTPS URLs from a given text string."""
    if not text:
        return []
    return URL_PATTERN.findall(text)


def decode_data_url(name: str, data_url: str) -> bytes:
    """Base64-decode an attachment, validating estimated payload size first."""
    if ";base64," in data_url:
        b64_content = data_url.split(";base64,", 1)[1]
    else:
        b64_content = data_url

    # 4 base64 chars per 3 bytes, ignoring the '=' padding tail.
    estimated_bytes = (len(b64_content) * 3) // 4
    if estimated_bytes > MAX_ATTACHMENT_BYTES:
        raise ValueError(
            f"附件 {name} 约 {estimated_bytes // (1024 * 1024)}MB，"
            f"超过 {MAX_ATTACHMENT_BYTES // (1024 * 1024)}MB 上限"
        )
    return base64.b64decode(b64_content)


def clean_html_text(html_content: str) -> str:
    """Strip HTML tags, styles, scripts and return cleaned readable text."""
    # 1. Remove script, style, head, noscript
    text = re.sub(r"<(script|style|head|noscript)[^>]*>.*?</\1>", "", html_content, flags=re.DOTALL | re.IGNORECASE)
    # 2. Convert common block tags to newlines
    text = re.sub(r"<(?:p|div|br|h[1-6]|li|tr)[^>]*>", "\n", text, flags=re.IGNORECASE)
    # 3. Strip remaining tags
    text = re.sub(r"<[^>]+>", "", text)
    # 4. Unescape HTML entities
    text = html.unescape(text)
    # 5. Normalize whitespace
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    return "\n".join(lines)


def parse_bytes(name: str, mime_type: str, raw_bytes: bytes, data_url: str = "") -> Dict[str, Any]:
    """
    Parse raw document/image bytes into standard payload dict:
    {
        "type": "image" | "document",
        "name": name,
        "mime_type": mime_type,
        "data_url": data_url,
        "extracted_text": str
    }
    """
    lower_name = name.lower()
    is_image = (
        mime_type.startswith("image/")
        or any(lower_name.endswith(ext) for ext in [".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp"])
    )

    if is_image:
        if not data_url:
            b64_content = base64.b64encode(raw_bytes).decode("ascii")
            data_url = f"data:{mime_type or 'image/jpeg'};base64,{b64_content}"
        return {
            "type": "image",
            "name": name,
            "mime_type": mime_type or "image/jpeg",
            "data_url": data_url,
            "extracted_text": ""
        }

    extracted_text = ""

    # 1. PDF Parsing
    if lower_name.endswith(".pdf") or "pdf" in mime_type:
        try:
            import pypdf
            reader = pypdf.PdfReader(io.BytesIO(raw_bytes))
            pages = []
            for idx, page in enumerate(reader.pages):
                txt = page.extract_text()
                if txt and txt.strip():
                    pages.append(f"--- 第 {idx + 1} 页 ---\n{txt.strip()}")
            extracted_text = "\n\n".join(pages)
        except ImportError:
            logger.warning("pypdf not available; falling back to raw stream text extraction")
            text_chunks = re.findall(rb"\(([^\)]+)\)Tj", raw_bytes)
            if text_chunks:
                extracted_text = "".join(c.decode("latin1", errors="ignore") for c in text_chunks)
            else:
                extracted_text = "[提示: 后端未安装 pypdf 库，建议安装 pypdf 以获得完整 PDF 解析效果]"
        except Exception as e:
            logger.warning("Failed to parse PDF %s: %s", name, e)
            extracted_text = f"[PDF 解析异常: {e}]"

    # 2. DOCX Word Document Parsing
    elif lower_name.endswith(".docx"):
        try:
            with zipfile.ZipFile(io.BytesIO(raw_bytes)) as z:
                xml_content = z.read("word/document.xml").decode("utf-8")
                chunks = re.findall(r"<w:t[^>]*>(.*?)</w:t>", xml_content)
                extracted_text = "".join(chunks)
        except Exception as e:
            logger.warning("Failed to parse docx %s: %s", name, e)
            extracted_text = f"[Word 文档解析失败: {e}]"

    # 3. Excel Spreadsheet (.xlsx) Parsing
    elif lower_name.endswith(".xlsx"):
        try:
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(raw_bytes), data_only=True)
            sheet_texts = []
            for sheetname in wb.sheetnames[:5]:
                ws = wb[sheetname]
                rows = []
                for row in ws.iter_rows(values_only=True):
                    row_vals = [str(v).strip() if v is not None else "" for v in row]
                    if any(row_vals):
                        rows.append(" | ".join(row_vals))
                if rows:
                    sheet_texts.append(f"### 工作表: {sheetname}\n" + "\n".join(rows[:200]))
            extracted_text = "\n\n".join(sheet_texts)
        except ImportError:
            # Fallback to standard library zipfile parsing of sheet XML
            try:
                with zipfile.ZipFile(io.BytesIO(raw_bytes)) as z:
                    shared_strings = []
                    if "xl/sharedStrings.xml" in z.namelist():
                        ss_xml = z.read("xl/sharedStrings.xml").decode("utf-8")
                        shared_strings = re.findall(r"<t[^>]*>(.*?)</t>", ss_xml)
                    sheets = [n for n in z.namelist() if n.startswith("xl/worksheets/sheet") and n.endswith(".xml")]
                    sheet_texts = []
                    for sheet_file in sheets[:3]:
                        s_xml = z.read(sheet_file).decode("utf-8")
                        rows_xml = re.findall(r"<row[^>]*>(.*?)</row>", s_xml)
                        rows = []
                        for r_xml in rows_xml[:100]:
                            cells = re.findall(r'<c[^>]*?(?: t="([^"]*)")?[^>]*>(?:<v>(.*?)</v>)?</c>', r_xml)
                            row_vals = []
                            for cell_type, val in cells:
                                if cell_type == "s" and val.isdigit() and int(val) < len(shared_strings):
                                    row_vals.append(shared_strings[int(val)])
                                else:
                                    row_vals.append(val or "")
                            if any(row_vals):
                                rows.append(" | ".join(row_vals))
                        if rows:
                            sheet_texts.append(f"### 表格: {sheet_file}\n" + "\n".join(rows))
                    extracted_text = "\n\n".join(sheet_texts) if sheet_texts else "[Excel 文件未包含可提取的数据内容]"
            except Exception as ze:
                logger.warning("Failed to parse xlsx %s: %s", name, ze)
                extracted_text = f"[Excel 表格解析失败: {ze}]"
        except Exception as e:
            logger.warning("Failed to parse xlsx %s: %s", name, e)
            extracted_text = f"[Excel 表格解析失败: {e}]"

    # 4. Plain Text / Code / Markdown / CSV / JSON
    else:
        for enc in ["utf-8", "gbk", "gb2312", "latin1"]:
            try:
                extracted_text = raw_bytes.decode(enc)
                break
            except UnicodeDecodeError:
                continue


    if not extracted_text:
        extracted_text = f"[文件 {name} 内容为空或无法识别为纯文本]"

    # Apply length guardrail
    if len(extracted_text) > MAX_DOCUMENT_CHARS:
        extracted_text = extracted_text[:MAX_DOCUMENT_CHARS] + f"\n\n[注意：文档过长，已截取前 {MAX_DOCUMENT_CHARS} 字符供分析]"

    return {
        "type": "document",
        "name": name,
        "mime_type": mime_type,
        "data_url": data_url,
        "extracted_text": extracted_text.strip()
    }


def parse_attachment(name: str, mime_type: str, data_url: str) -> Dict[str, Any]:
    """
    Parse an attachment payload into either an image payload or extracted text document.
    """
    is_image = (
        mime_type.startswith("image/")
        or any(name.lower().endswith(ext) for ext in [".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp"])
    )

    if is_image:
        return {
            "type": "image",
            "name": name,
            "mime_type": mime_type or "image/jpeg",
            "data_url": data_url,
            "extracted_text": ""
        }

    # Extract base64 payload
    try:
        raw_bytes = decode_data_url(name, data_url)
    except Exception as e:
        logger.error("Failed to decode base64 attachment %s: %s", name, e)
        return {
            "type": "document",
            "name": name,
            "mime_type": mime_type,
            "data_url": data_url,
            "extracted_text": f"[无法读取文件 {name}: {e}]"
        }

    return parse_bytes(name, mime_type, raw_bytes, data_url=data_url)


async def fetch_and_parse_url(url: str) -> Optional[Dict[str, Any]]:
    """
    Download and parse content from a remote URL (PDF, Word, Image, Webpage, etc.).
    Returns standard parsed dictionary or None on failure.
    """
    try:
        parsed_url = urllib.parse.urlparse(url)
        path = parsed_url.path
        filename = path.split("/")[-1] if "/" in path else ""
        if not filename:
            filename = parsed_url.netloc or "downloaded_document"

        try:
            fetched = await fetch_public_bytes(
                url,
                max_bytes=MAX_REMOTE_FETCH_BYTES,
                timeout=25.0,
                headers=_FETCH_HEADERS,
            )
        except ContentTooLargeError:
            return {
                "type": "document",
                "name": filename,
                "mime_type": "",
                "data_url": url,
                "extracted_text": f"[文件 {filename} 体积超过 25MB 限制，暂不支持自动提取]"
            }
        except UnsafeURLError as e:
            logger.warning("Refused to fetch non-public URL %s: %s", url, e)
            return None
        except httpx.HTTPStatusError as e:
            logger.warning("Failed to fetch remote URL %s: HTTP %s", url, e.response.status_code)
            return None

        content_type = fetched.content_type.lower()
        raw_bytes = fetched.data

        # Check if image
        is_image = (
            content_type.startswith("image/")
            or any(filename.lower().endswith(ext) for ext in [".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp"])
        )
        if is_image:
            b64_content = base64.b64encode(raw_bytes).decode("ascii")
            mime = content_type.split(";")[0] if content_type else "image/jpeg"
            return {
                "type": "image",
                "name": filename,
                "mime_type": mime,
                "data_url": f"data:{mime};base64,{b64_content}",
                "extracted_text": ""
            }

        # Check if HTML webpage (and not a direct document download)
        is_doc_ext = any(filename.lower().endswith(ext) for ext in [".pdf", ".docx", ".doc", ".txt", ".csv", ".json", ".md"])
        if "text/html" in content_type and not is_doc_ext:
            charset_match = _CHARSET_PATTERN.search(fetched.content_type)
            encoding = charset_match.group(1) if charset_match else "utf-8"
            try:
                html_text = raw_bytes.decode(encoding, errors="ignore")
            except LookupError:
                html_text = raw_bytes.decode("utf-8", errors="ignore")
            extracted_text = clean_html_text(html_text)
            if len(extracted_text) > MAX_DOCUMENT_CHARS:
                extracted_text = extracted_text[:MAX_DOCUMENT_CHARS] + f"\n\n[注意：网页过长，已截取前 {MAX_DOCUMENT_CHARS} 字符]"
            return {
                "type": "document",
                "name": filename or "网页内容",
                "mime_type": "text/html",
                "data_url": url,
                "extracted_text": extracted_text
            }

        # Parse as document bytes (PDF, Word, Text, etc.)
        return parse_bytes(name=filename, mime_type=content_type, raw_bytes=raw_bytes, data_url=url)

    except Exception as e:
        logger.warning("Error fetching or parsing remote URL %s: %s", url, e)
        return None
